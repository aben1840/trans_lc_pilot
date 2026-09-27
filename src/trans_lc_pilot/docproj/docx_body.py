"""Turn an HTML fragment into content for a docx body.

The mapping is best-effort by design: HTML can express things a docx will
not take verbatim, and the reverse. Whatever cannot be carried over is
reported as a warning rather than dropped in silence, so a caller can
tell "nothing was lost" from "something was".
"""
from __future__ import annotations

from base64 import b64decode
from io import BytesIO
from urllib.parse import unquote_to_bytes

from bs4 import BeautifulSoup, NavigableString, Tag
from docx.document import Document as DocxDocument
from docx.image.exceptions import UnrecognizedImageError
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .docx_styles import Styles
from .html_headings import is_significant

_HEADING_LEVELS = {f"h{level}": level for level in range(1, 7)}
_INLINE = {
    "strong": "bold",
    "b": "bold",
    "em": "italic",
    "i": "italic",
    "u": "underline",
}
_CONTAINERS = frozenset({"span", "div", "ul", "ol", "li", "tbody", "tr", "td", "th"})


def append_fragment(html: str, document: DocxDocument, styles: Styles) -> list[str]:
    """Append ``html`` to ``document`` as body content.

    Args:
        html: A standalone document — as a bundle writes its pieces — or
            a bare fragment; a ``<body>`` wrapper is unwrapped either
            way.
        document: The document to append to, typically a template whose
            body has been cleared.
        styles: Styles resolved from that template.

    Returns:
        list[str]: One message per construct that could not be carried
        over, empty when the fragment was represented in full.
    """
    soup = BeautifulSoup(html, "html.parser")
    root = soup.body if soup.body is not None else soup
    warnings: list[str] = []
    for node in root.children:
        if not is_significant(node):
            continue
        _append_block(node, document, styles, warnings)
    return warnings


def _append_block(
    node: Tag, document: DocxDocument, styles: Styles, warnings: list[str]
) -> None:
    """Append one top-level node to the document.

    Args:
        node: A top-level node of the fragment.
        document: The document to append to.
        styles: Styles resolved from the template.
        warnings: Collects what could not be carried over.
    """
    if not isinstance(node, Tag):
        document.add_paragraph(str(node), style=styles.body)
    elif node.name in _HEADING_LEVELS:
        level = _HEADING_LEVELS[node.name]
        style = styles.heading.get(level)
        if style is None:
            warnings.append(
                f"the template has no style for <{node.name}>;"
                " the body style was used"
            )
        paragraph = document.add_paragraph(style=style or styles.body)
        _add_runs(paragraph, node, {}, warnings)
    elif node.name == "p":
        _add_runs(document.add_paragraph(style=styles.body), node, {}, warnings)
    elif node.name == "table":
        _append_table(node, document, styles, warnings)
    elif node.name in {"ul", "ol"}:
        _append_list(node, document, styles, node.name == "ol", warnings)
    elif node.name == "img":
        _append_image(node, document.add_paragraph(style=styles.body), warnings)
    elif node.name in _CONTAINERS:
        warnings.append(f"<{node.name}> is flattened into a paragraph")
        _add_runs(document.add_paragraph(style=styles.body), node, {}, warnings)
    else:
        warnings.append(f"unsupported element dropped: <{node.name}>")


def _add_runs(
    paragraph: Paragraph, node: Tag, formats: dict[str, bool], warnings: list[str]
) -> None:
    """Add ``node``'s inline content to ``paragraph`` as runs.

    Args:
        paragraph: The paragraph to fill.
        node: A block or inline element whose children are walked.
        formats: Run attributes inherited from enclosing inline tags.
        warnings: Collects what could not be carried over.
    """
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if text and not _is_layout_whitespace(text):
                _add_run(paragraph, text, formats)
        elif isinstance(child, Tag):
            _add_inline(child, paragraph, formats, warnings)


def _add_inline(
    node: Tag, paragraph: Paragraph, formats: dict[str, bool], warnings: list[str]
) -> None:
    """Add one inline element to ``paragraph``.

    Args:
        node: The inline element.
        paragraph: The paragraph to fill.
        formats: Run attributes inherited from enclosing inline tags.
        warnings: Collects what could not be carried over.
    """
    name = node.name
    if name in _INLINE:
        _add_runs(paragraph, node, {**formats, _INLINE[name]: True}, warnings)
    elif name == "br":
        paragraph.add_run().add_break()
    elif name == "img":
        _append_image(node, paragraph, warnings)
    elif name == "a":
        warnings.append("hyperlink targets are dropped; the link text is kept")
        _add_runs(paragraph, node, formats, warnings)
    elif name in _CONTAINERS:
        _add_runs(paragraph, node, formats, warnings)
    else:
        warnings.append(f"<{name}> is kept as plain text")
        _add_runs(paragraph, node, formats, warnings)


def _add_run(paragraph: Paragraph, text: str, formats: dict[str, bool]) -> Run:
    """Add one run of text with the inherited formatting.

    Args:
        paragraph: The paragraph to fill.
        text: The run's text.
        formats: Run attributes, e.g. ``{"bold": True}``.

    Returns:
        Run: The added run.
    """
    run = paragraph.add_run(text)
    for attribute, value in formats.items():
        setattr(run, attribute, value)
    return run


def _append_table(
    node: Tag, document: DocxDocument, styles: Styles, warnings: list[str]
) -> None:
    """Append an HTML table as a docx table.

    Args:
        node: The ``<table>`` element.
        document: The document to append to.
        styles: Styles resolved from the template.
        warnings: Collects what could not be carried over.
    """
    rows = [row.find_all(["td", "th"]) for row in node.find_all("tr")]
    rows = [row for row in rows if row]
    if not rows:
        warnings.append("empty <table> dropped")
        return

    table = document.add_table(rows=len(rows), cols=max(_width(row) for row in rows))
    if styles.table is None:
        warnings.append("the template has no table style; the table is unstyled")
    else:
        table.style = styles.table

    columns = len(table.columns)
    for index, row in enumerate(rows):
        column = 0
        for cell in row:
            if _span(cell, "rowspan") > 1:
                warnings.append("rowspan is not carried over")
            span = _span(cell, "colspan")
            target = table.cell(index, min(column, columns - 1))
            if span > 1:
                last = min(column + span - 1, columns - 1)
                target = target.merge(table.cell(index, last))
            _add_runs(target.paragraphs[0], cell, {}, warnings)
            column += span

    if node.find("table") is not None:
        warnings.append("nested tables are flattened")


def _append_list(
    node, document: DocxDocument, styles: Styles, ordered: bool, warnings: list[str]
) -> None:
    """Append an HTML list as docx paragraphs.

    Args:
        node: The ``<ul>`` or ``<ol>`` element.
        document: The document to append to.
        styles: Styles resolved from the template.
        ordered: Whether the list is ordered.
        warnings: Collects what could not be carried over.
    """
    style = styles.number if ordered else styles.bullet
    if style is None:
        kind = "numbered" if ordered else "bullet"
        warnings.append(
            f"the template has no {kind} list style; the body style was used"
        )
    for item in node.find_all("li", recursive=False):
        if item.find(["ul", "ol"], recursive=False) is not None:
            warnings.append("nested lists are flattened")
        paragraph = document.add_paragraph(style=style or styles.body)
        _add_runs(paragraph, item, {}, warnings)


def _append_image(node: Tag, paragraph: Paragraph, warnings: list[str]) -> None:
    """Embed an inline image into ``paragraph``.

    Args:
        node: The ``<img>`` element.
        paragraph: The paragraph to append to.
        warnings: Collects what could not be carried over.
    """
    data = _image_bytes(node)
    if data is None:
        warnings.append(
            "image is not an inlined data URI and was dropped;"
            " only images carried inside the HTML survive"
        )
        return
    try:
        paragraph.add_run().add_picture(BytesIO(data))
    except (UnrecognizedImageError, ValueError, OSError) as exc:
        warnings.append(f"image could not be embedded: {exc}")


def _image_bytes(node: Tag) -> bytes | None:
    """Decode an ``<img>`` element's data URI.

    Args:
        node: The ``<img>`` element.

    Returns:
        bytes | None: The image bytes, or ``None`` when the source is not
        an inline data URI.
    """
    source = node.get("src", "")
    if not isinstance(source, str) or not source.startswith("data:"):
        return None
    header, separator, payload = source.partition(",")
    if not separator:
        return None
    if "base64" in header:
        try:
            return b64decode(payload)
        except ValueError:
            return None
    return unquote_to_bytes(payload)


def _span(cell: Tag, attribute: str) -> int:
    """Read a cell's span attribute.

    Args:
        cell: A ``<td>`` or ``<th>`` element.
        attribute: ``"colspan"`` or ``"rowspan"``.

    Returns:
        int: The span, at least ``1``.
    """
    try:
        return max(1, int(cell.get(attribute)))
    except (TypeError, ValueError):
        return 1


def _width(row: list[Tag]) -> int:
    """Return how many columns a table row occupies.

    Args:
        row: The cells of one row.

    Returns:
        int: The sum of the row's column spans.
    """
    return sum(_span(cell, "colspan") for cell in row)


def _is_layout_whitespace(text: str) -> bool:
    """Whether text is whitespace introduced by HTML formatting.

    A run of spaces between inline elements is content and is kept; the
    newlines a document is formatted with are not.

    Args:
        text: A text node's content.

    Returns:
        bool: ``True`` when the text is blank and spans lines.
    """
    return not text.strip() and "\n" in text
