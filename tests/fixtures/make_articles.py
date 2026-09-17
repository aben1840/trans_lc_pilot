"""Generate ``tests/fixtures/articles.docx``.

The fixture holds three short articles, separated by a ``Heading 1`` and a
blank paragraph. Articles 1 and 2 are Chinese; article 3 is English, so
the file exercises both a layout signal (heading style, blank paragraph)
and a content signal (language change).

Regenerate with::

    uv run python tests/fixtures/make_articles.py
"""
from __future__ import annotations

from pathlib import Path

import docx

OUT = Path(__file__).with_name("articles.docx")

ARTICLES: list[tuple[str, list[str]]] = [
    (
        "上海简介",
        [
            "上海位于中国东部沿海，地处长江入海口，是中国最大的城市之一。",
            "它是重要的国际金融、航运和贸易中心，外滩与陆家嘴隔黄浦江相望。",
            "全市常住人口约两千五百万，拥有众多高校、博物馆和美术馆。",
        ],
    ),
    (
        "北京简介",
        [
            "北京位于中国北部，是中华人民共和国的首都。",
            "它有三千年以上的建城史，故宫、天坛和长城都是世界闻名的古迹。",
            "北京也是全国的政治、文化和国际交往中心。",
        ],
    ),
    (
        "New York",
        [
            "New York City sits on the eastern coast of the United States, "
            "at the mouth of the Hudson River.",
            "It is the most populous city in the country and a global center "
            "for finance, culture, and media.",
            "Its five boroughs — Manhattan, Brooklyn, Queens, the Bronx, and "
            "Staten Island — hold over eight million residents.",
        ],
    ),
]


def main() -> None:
    """Build the fixture document."""
    document = docx.Document()
    for index, (title, paragraphs) in enumerate(ARTICLES):
        if index:
            document.add_paragraph("")
        document.add_heading(title, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
    document.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
