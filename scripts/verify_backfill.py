"""Verify backfill writer preserves OOXML formatting."""
from __future__ import annotations

import copy
import dataclasses
from pathlib import Path

from docx import Document

from trans_lc_pilot.docproj import read
from trans_lc_pilot.docproj.writers import write

proj = read("tests/fixtures/formatted.docx")

translated = copy.deepcopy(proj)
for b in translated.blocks:
    if b.kind in ("heading", "paragraph") and b.docx_para_idx is not None:
        b.text = "[TR] " + b.text
        if b.spans:
            b.spans = [
                dataclasses.replace(s, text="[TR] " + s.text)
                for s in b.spans
            ]

out = Path(".tmp/backfill_result.docx")
out.parent.mkdir(exist_ok=True)
write(translated, out, original_path=Path("tests/fixtures/formatted.docx"))
print(f"backfill written to {out}")

doc = Document(str(out))
print(f"paragraphs: {len(doc.paragraphs)}")
print(f"tables: {len(doc.tables)}")
print()

for i, p in enumerate(doc.paragraphs):
    runs_info = []
    for run in p.runs:
        flags = []
        if run.bold:
            flags.append("B")
        if run.italic:
            flags.append("I")
        if run.underline:
            flags.append("U")
        if run.font.color and run.font.color.rgb:
            flags.append(f"color:{run.font.color.rgb}")
        if run.font.size:
            flags.append(f"pt:{run.font.size.pt:.1f}")
        runs_info.append(f"[{run.text!r}|{','.join(flags)}]")
    print(f"  para[{i}] align={p.alignment} runs={' '.join(runs_info)}")